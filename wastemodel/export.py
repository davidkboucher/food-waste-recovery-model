"""Write results to one Excel workbook, color-coded with the same key as the discard log."""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Same colors as the Keys sheet in the source workbook
DONATABLE_COLORS = {
    "Yes": "C6EFCE", "No": "FFC7CE", "Maybe": "FFF696",
    "Weak Maybe": "FCD251", "Strong Maybe": "D1FC74",
}
MONEY, LBS, NUM = '"$"#,##0.00', '#,##0.0', '#,##0.##'

SHEETS = {
    "Donation Scenarios": ("scenarios", {
        "scenario": ("Scenario", None), "events": ("Discards", "#,##0"),
        "revenue_value_per_week": ("Menu Value / Week", MONEY),
        "food_cost_per_week": ("Food Cost / Week", MONEY),
        "deduction_per_week": ("Deduction / Week", MONEY),
        "tax_benefit_per_week": ("Tax Benefit / Week", MONEY),
        "lbs_per_week": ("Lbs / Week", LBS),
    }),
    "Discards (Valued)": ("events", {
        "date": ("Date", "m/d/yyyy"), "item": ("Food Item", None), "qty": ("Quantity (Units)", NUM),
        "reason": ("Discard Reason", None), "donatable": ("Donatable?", None),
        "conditions": ("Conditions", None), "tier": ("What It Would Take", None),
        "servings_lost": ("Servings Lost", NUM), "revenue_low": ("Value (Low)", MONEY),
        "revenue_mid": ("Value (Mid)", MONEY), "revenue_high": ("Value (High)", MONEY),
        "food_cost": ("Food Cost", MONEY), "lbs": ("Lbs", LBS),
    }),
    "Price Estimates": ("estimates", {
        "item": ("Food Item", None), "low": ("Low / Serving", MONEY),
        "mid": ("Mid / Serving", MONEY), "high": ("High / Serving", MONEY),
    }),
    "Data Quality": ("bad_refs", {
        "item": ("Food Item", None), "dish": ("Dish Listed", None), "problem": ("Problem", None),
    }),
}


def write_results(res, path):
    wb = Workbook()
    wb.remove(wb.active)
    for i, (title, (key, columns)) in enumerate(SHEETS.items(), start=1):
        df = res[key][list(columns)]
        ws = wb.create_sheet(title)
        ws.append([header for header, _ in columns.values()])
        for row in df.itertuples(index=False):
            ws.append([v.to_pydatetime() if hasattr(v, "to_pydatetime") else v for v in row])
        last_col = get_column_letter(len(columns))
        for j, (header, fmt) in enumerate(columns.values(), start=1):
            letter = get_column_letter(j)
            longest = max([len(header) + 4] + [len(str(v)) + 2 for v in df.iloc[:, j - 1].head(200)])
            ws.column_dimensions[letter].width = 12 if fmt == "m/d/yyyy" else min(40, max(12, longest))
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
        if len(df):
            table = Table(displayName=f"Table{i}", ref=f"A1:{last_col}{len(df) + 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        if key == "events":                    # color each row by its Donatable? label
            rng = f"A2:{last_col}{len(df) + 1}"
            for label, color in DONATABLE_COLORS.items():
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'$E2="{label}"'], fill=PatternFill(bgColor=color), stopIfTrue=True))
    for ws in wb:
        for cell in ws[1]:
            cell.font = Font(bold=True)
    wb.save(path)
